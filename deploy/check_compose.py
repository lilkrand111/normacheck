"""Run an isolated local Compose check; never uses the project's real storage.

Usage: .venv/bin/python deploy/check_compose.py --image normacheck:check
The image must already be built. Generated documents and diagnostics are retained
in a private temporary directory. Only this script's Compose project is stopped.
"""
import argparse
from concurrent.futures import ThreadPoolExecutor
import hashlib
from http.cookiejar import CookieJar
from io import BytesIO
import json
import os
from pathlib import Path
import secrets
import socket
import subprocess
import tempfile
import time
from urllib.error import HTTPError, URLError
from urllib.request import build_opener, HTTPCookieProcessor, Request

from docx import Document

ROOT = Path(__file__).resolve().parent.parent


def free_port():
    with socket.socket() as sock:
        sock.bind(("127.0.0.1", 0))
        return sock.getsockname()[1]


class Client:
    def __init__(self, url):
        self.url = url
        self.opener = build_opener(HTTPCookieProcessor(CookieJar()))
        status, _, body = self.request("/api/session")
        assert status == 200
        self.token = json.loads(body)["csrf_token"]

    def request(self, path, *, method="GET", data=None, headers=None):
        extra = dict(headers or {})
        if method == "POST":
            extra.setdefault("X-CSRF-Token", self.token)
        request = Request(self.url + path, data=data, headers=extra, method=method)
        try:
            response = self.opener.open(request, timeout=205)
        except HTTPError as error:
            response = error
        with response:
            return response.status, response.headers, response.read()

    def upload(self, data):
        boundary = secrets.token_hex(16)
        body = (f'--{boundary}\r\nContent-Disposition: form-data; name="file"; filename="check.docx"\r\n'
                'Content-Type: application/vnd.openxmlformats-officedocument.wordprocessingml.document\r\n\r\n').encode()
        body += data + f"\r\n--{boundary}--\r\n".encode()
        status, _, response = self.request("/api/documents", method="POST", data=body,
                                         headers={"Content-Type": f"multipart/form-data; boundary={boundary}"})
        assert status == 201, (status, response[:300])
        return json.loads(response)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--image", default="normacheck:check")
    parser.add_argument("--keep-running", action="store_true", help="Keep the isolated project for browser checks")
    args = parser.parse_args()
    work = Path(tempfile.mkdtemp(prefix="normacheck-compose-"))
    project = "nkcheck-" + secrets.token_hex(4)
    port, tls_port = free_port(), free_port()
    while tls_port == port:
        tls_port = free_port()
    # The parent is mode 0700; mount targets must be writable by UID 10001 in Linux.
    # These directories contain only generated check data, never user storage.
    for name in ("storage", "tmp", "caddy-data", "caddy-config"):
        (work / name).mkdir(mode=0o777)
        (work / name).chmod(0o777)
    env = {
        "DOMAIN": "localhost", "SITE_ADDRESS": "http://localhost", "COOKIE_SECURE": "false",
        "APP_SECRET": secrets.token_urlsafe(48), "APP_IMAGE": args.image,
        "STORAGE_PATH": str(work / "storage"), "HTTP_BIND": "127.0.0.1",
        "HTTP_PORT": str(port), "HTTPS_PORT": str(tls_port),
        "DOCKER_SUBNET": "172.30.77.0/29", "PROXY_IP": "172.30.77.2", "APP_IP": "172.30.77.3",
        # Functional checks need multiple mutations without waiting for refill.
        "RATE_LIMIT_BURST": "64", "RATE_LIMIT_PER_MINUTE": "12",
        "RATE_LIMIT_MAX_CLIENTS": "4096", "MAX_CONCURRENT_JOBS": "2",
        "MAX_UPLOAD_MB": "20", "MAX_STORAGE_MB": "2048", "WORKER_MEMORY_MB": "768",
        "PROCESSING_TIMEOUT_SECONDS": "60", "UPLOAD_TIMEOUT_SECONDS": "120",
        "UPLOAD_IDLE_TIMEOUT_SECONDS": "30", "APP_MEMORY_LIMIT": "640m",
        "APP_MEMORY_SWAP_LIMIT": "2560m", "APP_CPU_LIMIT": "0.9",
        "PROXY_MEMORY_LIMIT": "96m", "PROXY_MEMORY_SWAP_LIMIT": "192m",
        "MAX_REQUEST_BODY_BYTES": "22020096", "PROXY_UPLOAD_TIMEOUT": "130s",
        "PROXY_RESPONSE_TIMEOUT": "200s", "CADDY_IMAGE": "caddy:2-alpine",
    }
    env_file = work / "check.env"

    def write_env():
        env_file.write_text("".join(f"{key}={value}\n" for key, value in env.items()))
        env_file.chmod(0o600)

    write_env()
    override = work / "mounts.json"
    override.write_text(json.dumps({"services": {
        "app": {"volumes": [{"type": "bind", "source": str(work / "tmp"), "target": "/tmp/normacheck"}]},
        "caddy": {"volumes": [
            {"type": "bind", "source": str(work / "caddy-data"), "target": "/data"},
            {"type": "bind", "source": str(work / "caddy-config"), "target": "/config"},
        ]},
    }}))
    compose = ["docker", "compose", "--env-file", str(env_file), "-p", project,
               "-f", str(ROOT / "compose.yaml"), "-f", str(override)]
    log_file = work / "compose.log"

    def command(*arguments, capture=False):
        with log_file.open("a") as log:
            result = subprocess.run([*compose, *arguments], cwd=ROOT, check=True,
                                    stdout=subprocess.PIPE if capture else log, stderr=log, text=True,
                                    env={**os.environ, **env})
        return result.stdout

    url = f"http://localhost:{port}"

    def wait_http():
        for _ in range(60):
            try:
                if Client(url).request("/health")[0] == 200:
                    return
            except (URLError, AssertionError, ConnectionError):
                time.sleep(0.25)
        raise RuntimeError("Local proxy did not become ready")

    print(f"Isolated check directory: {work}", flush=True)
    passed = False
    try:
        command("config", "--quiet")
        command("up", "-d", "--no-build", "--wait", "--wait-timeout", "90")
        command("exec", "-T", "caddy", "caddy", "validate", "--config", "/etc/caddy/Caddyfile")
        wait_http()
        print("Compose, healthcheck and Caddy configuration: OK", flush=True)
        assert command("exec", "-T", "app", "id", "-u", capture=True).strip() == "10001"
        client = Client(url)
        status, headers, page = client.request("/")
        assert status == 200 and b"/assets/app.js" in page
        assert "script-src 'self'" in headers["Content-Security-Policy"]
        assert client.request("/assets/app.js")[0] == 200
        assert client.request("/storage/.session-secret")[0] == 404
        assert client.request("/", headers={"Host": "attacker.example"})[0] >= 400
        document = Document()
        document.add_heading("Введение", 1)
        for index in range(300):
            document.add_paragraph(f"Проверка контейнерного запуска. Абзац {index}.")
        stream = BytesIO()
        document.save(stream)
        data = stream.getvalue()
        (work / "generated.docx").write_bytes(data)
        initial = client.upload(data)
        document_id = initial["document_id"]
        # Inspect hashes inside the non-root app: permissions remain private on Linux.
        hash_code = ("import hashlib,pathlib; "
                     f"print(hashlib.sha256(pathlib.Path('/data/documents/{document_id}/original.docx').read_bytes()).hexdigest())")
        assert command("exec", "-T", "app", "python", "-c", hash_code, capture=True).strip() == hashlib.sha256(data).hexdigest()
        endpoint = f"/api/documents/{document_id}"
        fixed_status, _, fixed_body = client.request(endpoint + "/fix", method="POST", data=b"")
        assert fixed_status == 200, fixed_body[:300]
        fixed = json.loads(fixed_body)
        assert fixed["state"] == "fixed" and fixed["current_version"] == "corrected"
        downloaded = client.request(fixed["download_url"])
        assert downloaded[0] == 200
        Document(BytesIO(downloaded[2]))
        other = Client(url)
        assert other.request(endpoint)[0] == 404
        assert client.request(endpoint + "/fix", method="POST", data=b"", headers={"X-CSRF-Token": "wrong"})[0] == 403
        print("UI/API, original hash, check/fix/download, owner isolation and CSRF: OK", flush=True)

        def concurrent_upload(_):
            return Client(url).upload(data)

        with ThreadPoolExecutor(max_workers=2) as pool:
            results = list(pool.map(concurrent_upload, range(2)))
        assert len({item["document_id"] for item in results}) == 2
        print("Two concurrent DOCX uploads through Caddy: OK", flush=True)
        command("stats", "--no-stream")
        env["RATE_LIMIT_BURST"] = "1"
        write_env()
        command("up", "-d", "--no-deps", "--force-recreate", "--wait", "--wait-timeout", "90", "app")
        wait_http()
        assert client.request(endpoint)[0] == 200
        assert client.request(fixed["download_url"])[2] == downloaded[2]
        assert command("exec", "-T", "app", "python", "-c", hash_code, capture=True).strip() == hashlib.sha256(data).hexdigest()
        print("Container recreation preserves session, reports, original and corrected copy: OK", flush=True)
        assert client.request("/api/documents", method="POST", data=b"")[0] == 422
        status, headers, body = client.request("/api/documents", method="POST", data=b"",
                                             headers={"X-Forwarded-For": "203.0.113.99"})
        assert status == 429 and int(headers["Retry-After"]) > 0, (status, body[:300])
        assert client.request(endpoint)[0] == 200
        print("Rate limit before multipart, spoofed IP rejected, report still readable: OK", flush=True)
        print(f"All Compose checks passed. Evidence and generated data retained at {work}", flush=True)
        passed = True
        if args.keep_running:
            print(f"Browser check: {url}/?document={document_id}", flush=True)
            print("Stop with: " + " ".join([*compose, "down"]), flush=True)
    finally:
        if not args.keep_running or not passed:
            command("down", "--timeout", "220")


if __name__ == "__main__":
    main()
