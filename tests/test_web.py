"""Integration tests exercise the real DOCX modules and isolated worker processes."""
import asyncio
from concurrent.futures import ThreadPoolExecutor
from dataclasses import replace
from datetime import timedelta
from io import BytesIO
import json
import threading
from zipfile import ZipFile

from docx import Document
from docx.shared import Mm
from fastapi.testclient import TestClient
import pytest

from app.config import Settings
from app.main import create_app
from app.runner import ProcessingError
from app.storage import Storage, utc_now


@pytest.fixture
def settings(tmp_path):
    return Settings(storage_dir=tmp_path / "storage", allowed_hosts=("testserver", "localhost", "127.0.0.1"))


@pytest.fixture
def client(settings):
    with TestClient(create_app(settings)) as instance:
        instance.headers["X-CSRF-Token"] = instance.get("/api/session").json()["csrf_token"]
        yield instance


@pytest.fixture
def docx_bytes():
    document = Document()
    document.add_heading("Введение", 1)
    document.add_paragraph("Документ описывает разработку программного обеспечения и порядок проверки.")
    document.add_heading("1 Основная часть", 1)
    paragraph = document.add_paragraph("Система выполняет формальную проверку технической документации.")
    paragraph.paragraph_format.line_spacing = 1.5
    document.sections[0].page_width = Mm(216)
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def upload(client, data, filename="Пояснительная записка.docx"):
    return client.post("/api/documents", files={"file": (filename, data, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})


def test_complete_flow_preserves_original_and_rechecks(client, docx_bytes):
    response = upload(client, docx_bytes)
    assert response.status_code == 201, response.text
    data = response.json()
    document_id = data["document_id"]
    assert data["state"] == "checked" and data["current_version"] == "original"
    assert data["expires_at"] is None
    assert any(item["rule_id"] == "F01" for item in data["report"]["findings"])
    assert len(data["limitations"]) == 31
    assert "owner_id" not in data
    directory = client.app.state.service.storage.path(document_id)
    assert (directory / "original.docx").read_bytes() == docx_bytes
    assert client.get(f"/api/documents/{document_id}/download").status_code == 409
    fixed = client.post(f"/api/documents/{document_id}/fix")
    assert fixed.status_code == 200, fixed.text
    result = fixed.json()
    assert result["state"] == "fixed" and result["current_version"] == "corrected"
    download = client.get(result["download_url"])
    assert download.status_code == 200
    assert "attachment" in download.headers["content-disposition"]
    corrected = Document(BytesIO(download.content))
    assert abs(corrected.sections[0].page_width.mm - 210) < 0.1
    assert [p.text for p in corrected.paragraphs] == [p.text for p in Document(BytesIO(docx_bytes)).paragraphs]
    assert (directory / "original.docx").read_bytes() == docx_bytes
    assert result["report"] == json.loads((directory / "corrected-report.json").read_text())
    assert client.post(f"/api/documents/{document_id}/fix").json() == result
    preserved = client_files(directory)
    assert client.delete(f"/api/documents/{document_id}").status_code == 405
    assert client_files(directory) == preserved
    assert client.get(f"/api/documents/{document_id}").json() == result


def test_cross_session_access_and_csrf(client, docx_bytes):
    document_id = upload(client, docx_bytes).json()["document_id"]
    original_cookie = dict(client.cookies)
    original_token = client.headers["X-CSRF-Token"]
    client.cookies.clear()
    client.headers["X-CSRF-Token"] = client.get("/api/session").json()["csrf_token"]
    for method, suffix in (("GET", ""), ("GET", "/download"), ("POST", "/fix")):
        assert client.request(method, f"/api/documents/{document_id}{suffix}").status_code == 404
    assert client.delete(f"/api/documents/{document_id}").status_code == 405
    client.cookies.update(original_cookie)
    client.headers["X-CSRF-Token"] = "invalid"
    assert client.post(f"/api/documents/{document_id}/fix").status_code == 403
    assert client.delete(f"/api/documents/{document_id}").status_code == 405
    client.headers["X-CSRF-Token"] = original_token
    assert client.get(f"/api/documents/{document_id}").status_code == 200


@pytest.mark.parametrize("name,data,status", [("file.pdf", b"text", 415), ("file.docx", b"", 422), ("file.docx", b"not a zip", 422)])
def test_invalid_uploads(client, name, data, status):
    response = upload(client, data, name)
    assert response.status_code == status, response.text
    assert "error" in response.json()
    assert str(client.app.state.settings.storage_dir) not in response.text


def test_upload_limit_and_chunked_body(settings):
    with TestClient(create_app(replace(settings, max_upload_bytes=100))) as client:
        client.headers["X-CSRF-Token"] = client.get("/api/session").json()["csrf_token"]
        assert upload(client, b"a" * 101).status_code == 413
        assert upload(client, b"a" * (1024 * 1024 + 101)).status_code == 413
        assert not list(client.app.state.service.storage.documents.iterdir())


def test_missing_and_multiple_files(client, docx_bytes):
    assert client.post("/api/documents").status_code == 422
    response = client.post("/api/documents", files=[("file", ("1.docx", docx_bytes)), ("file", ("2.docx", docx_bytes))])
    assert response.status_code == 400


def test_failed_fix_keeps_original_report_and_hides_download(client, docx_bytes, monkeypatch):
    initial = upload(client, docx_bytes).json()
    document_id = initial["document_id"]

    async def fail(source, work, operation):
        (work / "corrected.docx").write_bytes(docx_bytes)
        raise ProcessingError("processing_failed", "Ошибка повторной проверки.")

    monkeypatch.setattr(client.app.state.service.runner, "run", fail)
    assert client.post(f"/api/documents/{document_id}/fix").status_code == 500
    current = client.get(f"/api/documents/{document_id}").json()
    assert current["state"] == "failed"
    assert current["report"] == initial["report"]
    assert current["current_version"] == "original" and current["download_url"] is None
    assert client.get(f"/api/documents/{document_id}/download").status_code == 409
    assert not (client.app.state.service.storage.path(document_id) / "corrected.docx").exists()


def test_two_documents_are_processed_concurrently_and_capacity_is_bounded(client, docx_bytes, monkeypatch):
    service = client.app.state.service
    original_run = service.runner.run
    started, release = threading.Event(), threading.Event()
    active = []

    async def barrier(source, work, operation):
        active.append(source)
        if len(active) == 2:
            started.set()
        assert await asyncio.to_thread(release.wait, 10)
        return await original_run(source, work, operation)

    monkeypatch.setattr(service.runner, "run", barrier)
    with ThreadPoolExecutor(max_workers=2) as pool:
        futures = [pool.submit(upload, client, docx_bytes, f"document-{i}.docx") for i in range(2)]
        try:
            assert started.wait(10), "Both independent documents must reach processing before either completes"
            assert len(set(active)) == 2
            assert client.get("/health").status_code == 200
            busy = upload(client, docx_bytes)
            assert busy.status_code == 503 and busy.headers["retry-after"] == "5"
        finally:
            release.set()
        results = [future.result(timeout=15) for future in futures]
    assert all(result.status_code == 201 for result in results)
    assert len({result.json()["document_id"] for result in results}) == 2


def test_same_document_cannot_be_fixed_twice(client, docx_bytes, monkeypatch):
    document_id = upload(client, docx_bytes).json()["document_id"]
    service = client.app.state.service
    original_run = service.runner.run
    started, release = threading.Event(), threading.Event()

    async def barrier(source, work, operation):
        started.set()
        assert await asyncio.to_thread(release.wait, 10)
        return await original_run(source, work, operation)

    monkeypatch.setattr(service.runner, "run", barrier)
    with ThreadPoolExecutor(max_workers=1) as pool:
        future = pool.submit(client.post, f"/api/documents/{document_id}/fix")
        try:
            assert started.wait(10)
            assert client.post(f"/api/documents/{document_id}/fix").status_code == 409
            assert client.delete(f"/api/documents/{document_id}").status_code == 405
            current = client.get(f"/api/documents/{document_id}").json()
            assert current["state"] == "fixing" and current["report"]
        finally:
            release.set()
        assert future.result(timeout=15).status_code == 200


def test_timeout_kills_worker(settings, docx_bytes):
    with TestClient(create_app(replace(settings, processing_timeout=0.001))) as client:
        client.headers["X-CSRF-Token"] = client.get("/api/session").json()["csrf_token"]
        assert upload(client, docx_bytes).status_code == 504
        assert not client.app.state.service.runner.processes
        directory = next(client.app.state.service.storage.documents.iterdir())
        metadata = json.loads((directory / "metadata.json").read_text())
        assert metadata["state"] == "failed"
        assert (directory / "original.docx").read_bytes() == docx_bytes
        assert not list(directory.glob(".work-*"))


def test_restart_persists_session_and_report_and_recovers_interrupted_job(settings, docx_bytes):
    with TestClient(create_app(settings)) as first:
        first.headers["X-CSRF-Token"] = first.get("/api/session").json()["csrf_token"]
        document_id = upload(first, docx_bytes).json()["document_id"]
        cookies = dict(first.cookies)
        metadata = first.app.state.service.storage.read(document_id)
        metadata["state"] = "fixing"
        first.app.state.service.storage.write(metadata)
    with TestClient(create_app(settings)) as second:
        second.cookies.update(cookies)
        response = second.get(f"/api/documents/{document_id}")
        assert response.status_code == 200
        assert response.json()["state"] == "failed"
        assert response.json()["current_version"] == "original"
        assert response.json()["last_error"]["code"] == "interrupted"


def test_permanent_storage_keeps_old_documents_available(client, docx_bytes):
    assert client.get("/api/session").json()["retention_hours"] is None
    data = upload(client, docx_bytes).json()
    document_id = data["document_id"]
    service = client.app.state.service
    metadata = service.storage.read(document_id)
    assert metadata["expires_at"] is None
    metadata["created_at"] = (utc_now() - timedelta(days=365)).isoformat()
    # Even an old, not-yet-normalized deadline must not cause a 410 or deletion.
    metadata["expires_at"] = (utc_now() - timedelta(days=364)).isoformat()
    service.storage.write(metadata)
    response = client.get(f"/api/documents/{document_id}")
    assert response.status_code == 200
    assert response.json()["expires_at"] is None
    assert client.post(f"/api/documents/{document_id}/fix").status_code == 200
    assert client.get(f"/api/documents/{document_id}/download").status_code == 200
    assert (service.storage.path(document_id) / "original.docx").read_bytes() == docx_bytes


def test_startup_converts_legacy_expired_uploads_to_permanent(settings, docx_bytes):
    with TestClient(create_app(settings)) as first:
        first.headers["X-CSRF-Token"] = first.get("/api/session").json()["csrf_token"]
        data = upload(first, docx_bytes).json()
        assert data["expires_at"] is None
        document_id = data["document_id"]
        assert first.post(f"/api/documents/{document_id}/fix").status_code == 200
        storage = first.app.state.service.storage
        directory = storage.path(document_id)
        preserved = {path.name: path.read_bytes() for path in directory.iterdir() if path.name != "metadata.json"}
        metadata = storage.read(document_id)
        metadata["expires_at"] = (utc_now() - timedelta(days=30)).isoformat()
        storage.write(metadata)
        cookies = dict(first.cookies)
    with TestClient(create_app(settings)) as second:
        second.cookies.update(cookies)
        response = second.get(f"/api/documents/{document_id}")
        assert response.status_code == 200
        assert response.json()["expires_at"] is None
        assert second.app.state.service.storage.read(document_id)["expires_at"] is None
        assert second.get(f"/api/documents/{document_id}/download").status_code == 200
        assert client_files(directory) == preserved


def client_files(directory):
    return {path.name: path.read_bytes() for path in directory.iterdir() if path.name != "metadata.json"}


@pytest.mark.parametrize("legacy_value", ["24", "-1", "invalid"])
def test_old_cleanup_settings_cannot_enable_deletion(settings, docx_bytes, monkeypatch, legacy_value):
    monkeypatch.setenv("RETENTION_HOURS", legacy_value)
    monkeypatch.setenv("CLEANUP_INTERVAL_SECONDS", "0.001")
    configured = Settings(storage_dir=settings.storage_dir, allowed_hosts=settings.allowed_hosts)
    with TestClient(create_app(configured)) as first:
        first.headers["X-CSRF-Token"] = first.get("/api/session").json()["csrf_token"]
        data = upload(first, docx_bytes).json()
        assert data["expires_at"] is None
        document_id = data["document_id"]
        storage = first.app.state.service.storage
        metadata = storage.read(document_id)
        metadata["expires_at"] = (utc_now() - timedelta(days=30)).isoformat()
        storage.write(metadata)
        cookies = dict(first.cookies)
    with TestClient(create_app(configured)) as second:
        second.cookies.update(cookies)
        assert second.get(f"/api/documents/{document_id}").status_code == 200
        assert (second.app.state.service.storage.path(document_id) / "original.docx").read_bytes() == docx_bytes


def test_invalid_docx_remains_stored_after_failed_check(client):
    payload = b"unreadable DOCX submitted for processing"
    assert upload(client, payload).status_code == 422
    storage = client.app.state.service.storage
    directory = next(storage.documents.iterdir())
    assert (directory / "original.docx").read_bytes() == payload
    assert storage.read(directory.name)["state"] == "failed"


def test_interrupted_upload_keeps_received_bytes(client):
    class InterruptedUpload:
        filename = "interrupted.docx"
        size = None

        def __init__(self):
            self.reads = 0

        async def read(self, amount):
            self.reads += 1
            if self.reads == 1:
                return b"already received bytes"
            raise ConnectionError("transfer interrupted")

    service = client.app.state.service
    with pytest.raises(ConnectionError):
        client.portal.call(service.upload, InterruptedUpload(), "test-owner")
    directory = next(service.storage.documents.iterdir())
    assert (directory / ".upload").read_bytes() == b"already received bytes"
    assert service.storage.read(directory.name)["state"] == "failed"
    assert not service.operations


def test_repeated_uploads_keep_each_original(client, docx_bytes):
    first = upload(client, docx_bytes).json()["document_id"]
    second = upload(client, docx_bytes).json()["document_id"]
    assert first != second
    storage = client.app.state.service.storage
    for document_id in (first, second):
        assert (storage.path(document_id) / "original.docx").read_bytes() == docx_bytes


def test_path_traversal_and_private_storage_are_unavailable(client):
    assert client.get("/api/documents/not-a-uuid").status_code == 404
    assert client.get("/storage/.session-secret").status_code == 404
    assert client.get("/static/../storage/.session-secret").status_code == 404
    assert client.get("/", headers={"Host": "attacker.example"}).status_code == 400


def test_frontend_is_served_with_csp_and_api_stays_available(client):
    page = client.get("/")
    assert page.status_code == 200
    assert 'lang="ru"' in page.text
    assert 'src="/assets/app.js"' in page.text
    assert "script-src 'self'" in page.headers["content-security-policy"]
    assert page.headers["cache-control"] == "no-cache"
    for asset in ("app.js", "report.mjs", "styles.css", "favicon.svg"):
        assert client.get(f"/assets/{asset}").status_code == 200
    for path in ("/assets/../app/security.py", "/assets/%2e%2e/storage/.session-secret", "/assets/index.html", "/storage/documents"):
        assert client.get(path).status_code == 404
    assert client.get("/static/app.js").status_code == 404
    schema = client.get("/openapi.json").json()
    assert all("delete" not in methods for methods in schema["paths"].values())
    assert "/api/documents" in schema["paths"]
    request_body = schema["paths"]["/api/documents"]["post"]["requestBody"]
    assert "file" in request_body["content"]["multipart/form-data"]["schema"]["properties"]


def test_duplicate_zip_entries_are_rejected(client, docx_bytes):
    output = BytesIO()
    with ZipFile(BytesIO(docx_bytes)) as original, ZipFile(output, "w") as modified:
        for entry in original.infolist():
            modified.writestr(entry, original.read(entry))
        with pytest.warns(UserWarning):
            modified.writestr("word/document.xml", original.read("word/document.xml"))
    assert upload(client, output.getvalue()).status_code == 422


def test_configurable_four_parallel_documents(settings, docx_bytes, monkeypatch):
    with TestClient(create_app(replace(settings, max_concurrent_jobs=4))) as client:
        client.headers["X-CSRF-Token"] = client.get("/api/session").json()["csrf_token"]
        assert client.get("/api/session").json()["max_concurrent_jobs"] == 4
        service = client.app.state.service
        original_run = service.runner.run
        started, release = threading.Event(), threading.Event()
        active = []

        async def barrier(source, work, operation):
            active.append(source)
            if len(active) == 4:
                started.set()
            assert await asyncio.to_thread(release.wait, 10)
            return await original_run(source, work, operation)

        monkeypatch.setattr(service.runner, "run", barrier)
        with ThreadPoolExecutor(max_workers=4) as pool:
            futures = [pool.submit(upload, client, docx_bytes) for _ in range(4)]
            try:
                assert started.wait(10)
                assert len(set(active)) == 4
            finally:
                release.set()
            assert all(future.result(timeout=15).status_code == 201 for future in futures)


def test_truncated_or_invalid_report_is_not_published(client, docx_bytes, monkeypatch):
    document_id = upload(client, docx_bytes).json()["document_id"]

    async def invalid_report(source, work, operation):
        (work / "corrected.docx").write_bytes(docx_bytes)
        (work / "report.json").write_text('{"standard":"СТП 01–2017","findings":', encoding="utf-8")

    monkeypatch.setattr(client.app.state.service.runner, "run", invalid_report)
    assert client.post(f"/api/documents/{document_id}/fix").status_code == 500
    current = client.get(f"/api/documents/{document_id}").json()
    assert current["current_version"] == "original"
    assert current["download_url"] is None
    assert current["report"]["findings"]


def test_storage_budget_rejects_upload_before_creating_files(settings, docx_bytes):
    with TestClient(create_app(replace(settings, max_storage_bytes=1024))) as client:
        client.headers["X-CSRF-Token"] = client.get("/api/session").json()["csrf_token"]
        assert upload(client, docx_bytes).status_code == 507
        assert not list(client.app.state.service.storage.documents.iterdir())


def test_second_api_cannot_recover_live_jobs(client):
    second_storage = Storage(client.app.state.settings.storage_dir)
    with pytest.raises(RuntimeError, match="MAX_CONCURRENT_JOBS"):
        with second_storage.instance_lock():
            pytest.fail("Another API must not enter the same storage")


def test_chunked_upload_limit_is_enforced_without_content_length(settings):
    async def exercise():
        application = create_app(replace(settings, max_upload_bytes=100))
        async with application.router.lifespan_context(application):
            messages = []
            request_count = 0

            async def receive():
                nonlocal request_count
                request_count += 1
                return {"type": "http.request", "body": b"x" * (1024 * 1024 + 101), "more_body": False}

            async def send(message):
                messages.append(message)

            # Obtain a valid session directly from its serializer for this raw
            # ASGI request; no Content-Length header is supplied.
            sessions = application.state.sessions
            import secrets
            owner = secrets.token_urlsafe(32)
            cookie = "docx_session=" + sessions.serializer.dumps(owner)
            scope = {"type": "http", "asgi": {"version": "3.0"}, "http_version": "1.1",
                     "method": "POST", "scheme": "http", "path": "/api/documents", "raw_path": b"/api/documents", "query_string": b"",
                     "headers": [(b"host", b"testserver"), (b"content-type", b"multipart/form-data; boundary=test"),
                                 (b"cookie", cookie.encode()), (b"x-csrf-token", sessions.token(owner).encode())],
                     "client": ("127.0.0.1", 123), "server": ("testserver", 80)}
            await application(scope, receive, send)
            assert next(message for message in messages if message["type"] == "http.response.start")["status"] == 413
            assert request_count == 1
            assert not list(application.state.service.storage.documents.iterdir())

    asyncio.run(exercise())
